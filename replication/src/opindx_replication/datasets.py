"""Loss-checked public annual formats from the canonical website Parquets."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import shutil

import pandas as pd
import pyarrow.parquet as pq

def digest(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()

def description(name):
    known = {
        'openalex_id':'OpenAlex source identifier for the journal.', 'title':'Journal title.',
        'publisher':'Publisher from OpenAlex.', 'issn_l':'Linking ISSN; retain as text.',
        'issns':'Journal ISSNs; retain as text.', 'oa_domain':'Parent domain of the predominant OpenAlex field.',
        'oa_field':'Most frequent primary-topic field in classified articles/reviews in years t-5 through t-1.',
        'norwegian_area':'Norwegian Register area.', 'norwegian_field':'Norwegian Register assigned field(s).',
        'norwegian_level':'Norwegian Register annual level(s); retain as text.',
        'norwegian_register_url':'Register journal URLs; multiple IDs are separated by semicolons.',
        'is_open_access':'OpenAlex journal open-access status.', 'score_year':'Citation year t.',
        'reference_coverage_pct':'100 times Filtered / Raw publications; percentage points (0-100).',
        'active_years':'Number of years with Raw articles/reviews in the preceding five-year window.',
        'in_n':'Member of the Norwegian Register calculation universe in this year.',
        'in_oa':'Member of the OpenAlex calculation universe.',
        'oa_field_modal_share':'Predominant field share among classified articles/reviews, fraction 0-1.',
        'oa_field_modal_works':'Number of articles/reviews in the predominant field.',
        'oa_field_classified_works':'Number of articles/reviews with a named field in the five-year window.',
        'oa_field_classification_coverage':'Classified / all Raw articles/reviews, fraction 0-1.',
        'oa_field_tied_modes':'Number of fields tied for highest count; alphabetical tie resolution.',
        'oa_field_other_works':'Classified articles/reviews outside the top three fields.',
        'norwegian_entries_json':'For multiple register IDs: JSON array of ID, assigned field and title; empty for single IDs.',
        'norwegian_primary_id':'Unique title-matched primary register ID among multiple entries; otherwise empty.',
    }
    if name in known: return known[name]
    if name.startswith(('share_', 'per_article_')):
        metric = 'NF (Network Factor), normalized to sum to 100' if name.startswith('share_') else 'ANS (Article Network Score), article-weighted mean 1'
        universe = 'Norwegian Register' if '_n_' in name else 'OpenAlex'
        treatment = name.rsplit('_',1)[1].capitalize()
        return f'{metric}; {universe} universe; {treatment} treatment. Missing is distinct from zero.'
    if name.startswith('publications_'): return 'Articles/reviews in t-5 through t-1; '+name.rsplit('_',1)[1]+' treatment.'
    if name.startswith(('citations_n_', 'citations_oa_')):
        universe = 'Norwegian Register' if name.startswith('citations_n_') else 'OpenAlex'
        return f'Incoming citations used by NF/ANS in the {universe} universe in t, to articles/reviews in t-5 through t-1; {name.rsplit("_",1)[1]} treatment; journal self-citations excluded. Null outside the universe; zero is a measured zero.'
    if name.startswith('citations_'): return 'Recorded incoming citations from all cached citing sources in t to the five-year publication window, excluding journal self-citations; '+name.rsplit('_',1)[1]+' treatment.'
    if name.startswith('oa_field_') and name[-1:] in ['1','2','3']: return 'Field name at descending frequency rank '+name[-1]+'.'
    if name.startswith('oa_field_') and name.endswith('_works'): return 'Classified articles/reviews in '+name.removesuffix('_works')+'.'
    raise ValueError('Undocumented field: '+name)

def validate_values(expected, observed, *, spreadsheet=False):
    if expected is None:
        assert observed is None
    elif isinstance(expected, str):
        # Excel represents an empty string as an empty cell; identifiers stay strings.
        assert observed == expected or (spreadsheet and expected == '' and observed is None)
    elif isinstance(expected, bool):
        assert type(observed) is bool and observed == expected
    elif isinstance(expected, int):
        assert not isinstance(observed, bool) and observed == expected
    else:
        assert isinstance(observed, (int,float)) and not isinstance(observed,bool)
        assert math.isclose(expected,observed,rel_tol=2e-15,abs_tol=0), (expected,observed)

def export(source, output):
    import xlsxwriter
    from openpyxl import load_workbook
    source,output=Path(source),Path(output)
    output.mkdir(parents=True,exist_ok=True)
    manifest=json.loads((source/'manifest.json').read_text(encoding='utf-8'))
    reports=[]
    schema=None
    for year in manifest['years']:
        original=source/f'scores_{year}.parquet'
        expected_hash=manifest.get('assets',{}).get(original.name,{}).get('sha256')
        if expected_hash: assert digest(original)==expected_hash
        table=pq.read_table(original)
        if schema is None: schema=table.schema
        assert table.schema.equals(schema,check_metadata=False)
        frame=table.to_pandas()
        assert not frame.duplicated(['openalex_id','score_year']).any()
        assert frame.score_year.eq(year).all()
        metrics=[c for c in frame if c.startswith(('share_','per_article_'))]
        assert not frame[metrics].isna().all(axis=1).any()
        rows=table.to_pylist();names=table.column_names
        target=output/original.name
        if target.exists(): assert digest(target)==digest(original)
        else: shutil.copyfile(original,target)
        csv_path=output/f'scores_{year}.csv'
        with csv_path.open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f,quoting=csv.QUOTE_MINIMAL)
            writer.writerow(names)
            for row in rows:
                values=[row[c] for c in names]
                assert '\\N' not in values
                writer.writerow(['\\N' if v is None else str(v).lower() if isinstance(v,bool) else v for v in values])
        xlsx_path=output/f'scores_{year}.xlsx'
        with xlsxwriter.Workbook(xlsx_path,{'constant_memory':True,'strings_to_formulas':False,'strings_to_urls':False}) as book:
            sheet=book.add_worksheet('Journal metrics')
            header=book.add_format({'bold':True,'bg_color':'#E8EEF6','text_wrap':True,'valign':'top'})
            integer=book.add_format({'num_format':'#,##0'})
            numeric=book.add_format({'num_format':'0.###############'})
            sheet.freeze_panes(1,2);sheet.set_row(0,46)
            for i,name in enumerate(names):
                sheet.write_string(0,i,name,header)
                sheet.set_column(i,i,44 if name=='title' else 25)
            for i,row in enumerate(rows,1):
                for j,name in enumerate(names):
                    value=row[name]
                    if value is None: continue
                    if isinstance(value,str):
                        if len(value)>32767: raise ValueError('Excel cell length exceeded')
                        sheet.write_string(i,j,value)
                    elif isinstance(value,bool): sheet.write_boolean(i,j,value)
                    else: sheet.write_number(i,j,value,integer if isinstance(value,int) else numeric)
            sheet.autofilter(0,0,len(rows),len(names)-1)
            notes=book.add_worksheet('Read me')
            notes.set_column(0,0,110)
            for i,text in enumerate([
                f'Opindx complete annual dataset: citation year {year}',
                'NF = share_* columns; ANS = per_article_* columns. Full definitions are in Data dictionary.',
                'Raw includes all eligible publications; Filtered requires recorded references.',
                'n = Norwegian Register; oa = OpenAlex. Publication window: t-5 through t-1.',
                'Source snapshot: '+manifest['openalex_snapshot']+'; Norwegian Register: '+manifest['norwegian_register_snapshot'],
                'CC BY 4.0. OpenAlex and Norwegian Register numerical inputs; documented identity corrections.',
                'Missing values are blank; zero is a defined value. Excel numeric precision is approximately 15 digits.',
                'For exact double-precision values and the distinction between null and empty text, use Parquet or CSV.',
                'CSV null marker: \\N; CSV is UTF-8 with BOM; ISSNs and IDs must be imported as text.',
                'Percentiles depend on selected settings and are computed by the archived website engine.',
                'https://github.com/bread-n-roses/Opindx',
                'https://openalex.org', 'https://kanalregister.hkdir.no/en',
                'https://creativecommons.org/licenses/by/4.0/',
            ]): notes.write_string(i,0,text)
            dictionary=book.add_worksheet('Data dictionary')
            dictionary.set_column(0,0,35);dictionary.set_column(1,1,18);dictionary.set_column(2,2,100)
            dictionary.write_row(0,0,['Column','Type','Definition'],header)
            dictionary.freeze_panes(1,0)
            for i,field in enumerate(schema,1): dictionary.write_row(i,0,[field.name,str(field.type),description(field.name)])
        with csv_path.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.reader(f);assert next(reader)==names
            count=0
            for expected,actual in zip(rows,reader,strict=True):
                for c,value in zip(names,actual,strict=True):
                    e=expected[c]
                    if isinstance(e,bool): assert value in ('true','false')
                    parsed=None if value=='\\N' else value if isinstance(e,str) else value=='true' if isinstance(e,bool) else int(value) if isinstance(e,int) else float(value)
                    validate_values(e,parsed)
                    if isinstance(e,float): assert e==parsed
                count+=1
            assert count==len(rows)
        book=load_workbook(xlsx_path,read_only=True,data_only=False)
        sheet=book['Journal metrics'];iterator=sheet.iter_rows()
        assert [c.value for c in next(iterator)]==names
        count=0
        for expected,actual in zip(rows,iterator,strict=True):
            for name,cell in zip(names,actual,strict=True):
                assert cell.data_type!='f'
                validate_values(expected[name],cell.value,spreadsheet=True)
            count+=1
        book.close();assert count==len(rows)
        report={'year':year,'rows':len(rows),'columns':len(names),'csv_all_cells_exact':True,'xlsx_all_cells_verified':True,
                'xlsx_float_relative_tolerance':2e-15,'xlsx_empty_string_equals_blank':True,
                'files':{p.name:{'sha256':digest(p),'bytes':p.stat().st_size} for p in [target,csv_path,xlsx_path]}}
        reports.append(report)
        (output/'validation.json').write_text(json.dumps({'status':'IN_PROGRESS','annual':reports},indent=2)+'\n',encoding='utf-8')
        print(f'{year}: {len(rows):,} rows, {len(names)} columns; CSV and XLSX all-cell verification PASS',flush=True)
    with (output/'data_dictionary.csv').open('w',encoding='utf-8',newline='') as f:
        writer=csv.writer(f);writer.writerow(['column','type','definition'])
        for field in schema:writer.writerow([field.name,str(field.type),description(field.name)])
    shutil.copyfile(source/'manifest.json',output/'source_manifest.json')
    (output/'validation.json').write_text(json.dumps({'status':'PASS','annual':reports,'total_rows':sum(r['rows'] for r in reports)},indent=2)+'\n',encoding='utf-8')
    return reports

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path);parser.add_argument('output',type=Path)
    args=parser.parse_args();export(args.source,args.output)
